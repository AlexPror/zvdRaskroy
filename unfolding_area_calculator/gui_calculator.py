#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
GUI для расчета площадей разверток
Для цеха - универсальная утилита
"""

import sys
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog
from pathlib import Path
from datetime import datetime

from area_calculator import UnfoldingAreaCalculator


class UnfoldingAreaGUI:
    """GUI для расчета площадей"""
    
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("ZVD - Расчет площадей разверток v1.0")
        
        # Адаптивный размер
        screen_width = self.root.winfo_screenwidth()
        screen_height = self.root.winfo_screenheight()
        
        window_width = min(1050, int(screen_width * 0.75))
        window_height = min(750, int(screen_height * 0.8))
        
        x = (screen_width - window_width) // 2
        y = (screen_height - window_height) // 2
        
        self.root.geometry(f"{window_width}x{window_height}+{x}+{y}")
        self.root.minsize(850, 600)
        
        # Настройка стилей
        self.setup_styles()
        
        self.folder_path = tk.StringVar()
        self.order_number = tk.StringVar()
        
        self.files_data = []  # Список файлов с количествами
        self.calculator = UnfoldingAreaCalculator()
        
        self.create_widgets()
    
    def setup_styles(self):
        """Настройка красивых стилей"""
        style = ttk.Style()
        
        # Тема
        try:
            style.theme_use('clam')  # Современная тема
        except:
            pass
        
        # Цветовая схема ZVD
        ZVD_BLUE = '#1976D2'
        ZVD_LIGHT_BLUE = '#64B5F6'
        ZVD_DARK = '#0D47A1'
        ZVD_ACCENT = '#FF6F00'
        
        # Стиль для кнопок
        style.configure('Accent.TButton',
                       font=('Arial', 10, 'bold'),
                       foreground=ZVD_BLUE)
        
        # Стиль для заголовков
        style.configure('Title.TLabel',
                       font=('Arial', 16, 'bold'),
                       foreground=ZVD_DARK)
        
        style.configure('Subtitle.TLabel',
                       font=('Arial', 10),
                       foreground='#666666')
    
    def create_widgets(self):
        """Создание интерфейса"""
        
        # Основной контейнер
        main_frame = ttk.Frame(self.root, padding="15")
        main_frame.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        
        self.root.columnconfigure(0, weight=1)
        self.root.rowconfigure(0, weight=1)
        main_frame.columnconfigure(0, weight=1)
        # Вес для строк - чтобы таблица растягивалась
        main_frame.rowconfigure(4, weight=1)  # Строка с таблицей файлов
        
        # === ШАПКА: Логотип и название ===
        header_frame = ttk.Frame(main_frame)
        header_frame.grid(row=0, column=0, sticky=(tk.W, tk.E), pady=(0, 15))
        header_frame.columnconfigure(0, weight=1)
        
        # Заголовок
        title_label = ttk.Label(header_frame, 
                               text="📐 ZVD - Расчет площадей разверток", 
                               style='Title.TLabel')
        title_label.pack()
        
        subtitle_label = ttk.Label(header_frame,
                                  text="Универсальная утилита для производства",
                                  style='Subtitle.TLabel')
        subtitle_label.pack()
        
        # Разделитель
        ttk.Separator(main_frame, orient='horizontal').grid(row=1, column=0, sticky=(tk.W, tk.E), pady=10)
        
        # === СЕКЦИЯ 1: Выбор папки ===
        folder_frame = ttk.LabelFrame(main_frame, text="📁 Исходные данные", padding="10")
        folder_frame.grid(row=2, column=0, sticky=(tk.W, tk.E), pady=(0,10))
        folder_frame.columnconfigure(1, weight=1)
        
        ttk.Label(folder_frame, text="Папка с DXF:", 
                 font=('Arial', 9, 'bold')).grid(row=0, column=0, sticky=tk.W)
        ttk.Entry(folder_frame, textvariable=self.folder_path, 
                 font=('Arial', 9)).grid(row=0, column=1, sticky=(tk.W, tk.E), padx=5)
        ttk.Button(folder_frame, text="📁", command=self.browse_folder, width=3).grid(row=0, column=2)
        
        ttk.Label(folder_frame, text="Номер заказа:", 
                 font=('Arial', 9)).grid(row=1, column=0, sticky=tk.W, pady=(10,0))
        ttk.Entry(folder_frame, textvariable=self.order_number, 
                 font=('Arial', 9)).grid(row=1, column=1, sticky=(tk.W, tk.E), padx=5, pady=(10,0))
        
        # Большая кнопка загрузки
        load_btn = tk.Button(folder_frame, text="🔍 Загрузить файлы и рассчитать", 
                            command=self.load_files,
                            font=('Arial', 10, 'bold'),
                            bg='#1976D2',
                            fg='white',
                            activebackground='#0D47A1',
                            activeforeground='white',
                            relief=tk.RAISED,
                            bd=3,
                            cursor='hand2',
                            pady=8)
        load_btn.grid(row=2, column=0, columnspan=3, sticky=(tk.W, tk.E), pady=(15,5))
        
        # === СЕКЦИЯ 2: Список файлов с количествами ===
        files_frame = ttk.LabelFrame(main_frame, text="📋 Детали — введите количество в поле", padding="12")
        files_frame.grid(row=4, column=0, sticky=(tk.W, tk.E, tk.N, tk.S), pady=5)
        files_frame.columnconfigure(0, weight=1)
        files_frame.rowconfigure(0, weight=1)
        
        # Canvas для скроллинга таблицы (без фиксированной высоты)
        canvas_table = tk.Canvas(files_frame, highlightthickness=0)
        scrollbar_table = ttk.Scrollbar(files_frame, orient="vertical", command=canvas_table.yview)
        
        self.table_frame = ttk.Frame(canvas_table)
        self.table_frame.bind(
            "<Configure>",
            lambda e: canvas_table.configure(scrollregion=canvas_table.bbox("all"))
        )
        
        canvas_table.create_window((0, 0), window=self.table_frame, anchor="nw")
        canvas_table.configure(yscrollcommand=scrollbar_table.set)
        
        canvas_table.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        scrollbar_table.grid(row=0, column=1, sticky=(tk.N, tk.S))
        
        # Заголовки таблицы
        header_frame = tk.Frame(self.table_frame, bg='#1976D2')
        header_frame.grid(row=0, column=0, sticky=(tk.W, tk.E))
        
        # Ширины колонок (в символах) - ОДИНАКОВЫЕ для заголовков и данных!
        self.col_widths = {
            'num': 4,
            'file': 45,
            'width': 10,
            'height': 10,
            'area': 11,
            'qty': 9,
            'total': 11
        }
        
        # Заголовки с красивым фоном
        headers_config = [
            ("№", self.col_widths['num']),
            ("Файл развертки", self.col_widths['file']),
            ("Ширина\nмм", self.col_widths['width']),
            ("Высота\nмм", self.col_widths['height']),
            ("Площадь\nм²", self.col_widths['area']),
            ("Кол-во\nшт", self.col_widths['qty']),
            ("Итого\nм²", self.col_widths['total'])
        ]
        
        for col, (text, width) in enumerate(headers_config):
            label = tk.Label(header_frame, text=text, width=width, 
                           font=('Arial', 9, 'bold'),
                           bg='#1976D2', fg='white',
                           relief=tk.FLAT,
                           pady=6)
            label.grid(row=0, column=col, padx=0, sticky=(tk.W, tk.E))
        
        # Контейнер для строк данных
        self.data_rows_frame = ttk.Frame(self.table_frame)
        self.data_rows_frame.grid(row=1, column=0, sticky=(tk.W, tk.E))
        
        self.quantity_entries = []  # Список полей ввода количества
        
        # Скроллинг колесиком мыши
        def _on_mousewheel(event):
            canvas_table.yview_scroll(int(-1*(event.delta/120)), "units")
        canvas_table.bind_all("<MouseWheel>", _on_mousewheel)
        
        # Кнопки под таблицей
        buttons_frame = ttk.Frame(files_frame)
        buttons_frame.grid(row=1, column=0, columnspan=2, pady=(10,0))
        
        ttk.Button(buttons_frame, text="🔢 Одинаковое количество для всех", 
                  command=self.set_all_quantities,
                  width=32).pack(side=tk.LEFT, padx=5)
        ttk.Button(buttons_frame, text="🔄 Пересчитать итоги", 
                  command=self.recalculate_totals,
                  width=20).pack(side=tk.LEFT, padx=5)
        
        # === СЕКЦИЯ 3: Итоги и экспорт ===
        result_frame = ttk.LabelFrame(main_frame, text="📊 Результаты", padding="10")
        result_frame.grid(row=5, column=0, sticky=(tk.W, tk.E), pady=(10,5))
        result_frame.columnconfigure(0, weight=1)
        
        # Итоги с красивым оформлением
        totals_frame = ttk.Frame(result_frame)
        totals_frame.pack(pady=8, fill=tk.X)
        
        self.total_label = tk.Label(totals_frame, text="", 
                                    font=('Arial', 11, 'bold'),
                                    fg='#1976D2',
                                    bg='#E3F2FD',
                                    pady=8,
                                    relief=tk.RAISED,
                                    bd=2)
        self.total_label.pack(fill=tk.X, padx=5)
        
        # Кнопки экспорта - крупнее и ярче
        export_frame = ttk.Frame(result_frame)
        export_frame.pack(pady=8)
        
        ttk.Button(export_frame, text="📄 Экспорт в PDF", 
                  command=self.export_pdf, width=22,
                  style='Accent.TButton').pack(side=tk.LEFT, padx=8)
        ttk.Button(export_frame, text="📊 Экспорт в Excel", 
                  command=self.export_excel, width=22,
                  style='Accent.TButton').pack(side=tk.LEFT, padx=8)
        
        # Статусбар
        self.status_var = tk.StringVar(value="Готов к работе | Выберите папку с развертками")
        status_bar = ttk.Label(self.root, textvariable=self.status_var, 
                              relief=tk.SUNKEN, anchor=tk.W,
                              font=('Arial', 9))
        status_bar.grid(row=1, column=0, sticky=(tk.W, tk.E), padx=0, pady=0)
    
    def browse_folder(self):
        """Выбор папки"""
        folder = filedialog.askdirectory(title="Выберите папку с развертками")
        if folder:
            self.folder_path.set(folder)
    
    def load_files(self):
        """Загрузка файлов из папки"""
        folder = self.folder_path.get()
        
        if not folder or not Path(folder).exists():
            messagebox.showerror("Ошибка", "Выберите папку с файлами!")
            return
        
        self.status_var.set("Загрузка файлов...")
        
        # Очищаем старые строки
        for widget in self.data_rows_frame.winfo_children():
            widget.destroy()
        
        self.files_data = []
        self.quantity_entries = []
        
        # Ищем файлы - ТОЛЬКО DXF (универсальный формат)
        folder_path = Path(folder)
        files = list(folder_path.glob('*.dxf'))
        
        if not files:
            messagebox.showwarning("Внимание", 
                                 "Файлы DXF не найдены!\n\n"
                                 "Экспортируйте развертки в формат DXF:\n"
                                 "• SolidWorks → Файл → Сохранить как → DXF\n"
                                 "• КОМПАС → Файл → Сохранить как → DXF")
            self.status_var.set("Файлы DXF не найдены")
            return
        
        # Обрабатываем каждый файл
        for idx, file in enumerate(files):
            # Пытаемся извлечь количество из имени
            qty = 1
            filename_lower = file.stem.lower()
            
            import re
            match = re.search(r'(\d+)\s*шт', filename_lower)
            if match:
                qty = int(match.group(1))
            
            result = self.calculator.calculate_file(str(file), qty)
            
            if result['success']:
                self.files_data.append(result)
                
                # Создаем строку с виджетами
                row_frame = ttk.Frame(self.data_rows_frame)
                row_frame.grid(row=idx, column=0, sticky=(tk.W, tk.E), pady=1)
                
                # Чередуем цвет фона
                bg_color = '#F5F5F5' if idx % 2 == 0 else '#FFFFFF'
                
                # Используем те же ширины что и в заголовках!
                # Номер
                tk.Label(row_frame, text=str(idx+1), width=self.col_widths['num'], anchor='center',
                        bg=bg_color, font=('Arial', 9)).grid(row=0, column=0, padx=0)
                
                # Файл
                filename_text = result['filename']
                if len(filename_text) > 43:
                    filename_text = filename_text[:40] + "..."
                
                tk.Label(row_frame, text=filename_text, width=self.col_widths['file'], anchor='w',
                        bg=bg_color, font=('Arial', 9)).grid(row=0, column=1, padx=0, sticky=tk.W)
                
                # Габариты
                tk.Label(row_frame, text=f"{result['width']:.0f}", width=self.col_widths['width'], anchor='center',
                        bg=bg_color, font=('Arial', 9)).grid(row=0, column=2, padx=0)
                tk.Label(row_frame, text=f"{result['height']:.0f}", width=self.col_widths['height'], anchor='center',
                        bg=bg_color, font=('Arial', 9)).grid(row=0, column=3, padx=0)
                
                # Площадь
                tk.Label(row_frame, text=f"{result['area_m2']:.4f}", width=self.col_widths['area'], anchor='center',
                        bg=bg_color, font=('Arial', 9)).grid(row=0, column=4, padx=0)
                
                # ПОЛЕ ВВОДА КОЛИЧЕСТВА - выделяется цветом
                qty_var = tk.StringVar(value=str(qty))
                qty_entry = tk.Entry(row_frame, textvariable=qty_var, width=self.col_widths['qty'], 
                                    justify='center',
                                    font=('Arial', 10, 'bold'),
                                    bg='#FFF9C4',  # Желтый фон
                                    fg='#000000',
                                    relief=tk.SOLID,
                                    bd=2)
                qty_entry.grid(row=0, column=5, padx=0, pady=2)
                
                # Привязываем обновление при изменении
                qty_var.trace_add('write', lambda *args, idx=idx: self.on_quantity_changed(idx))
                
                self.quantity_entries.append({
                    'var': qty_var,
                    'entry': qty_entry,
                    'index': idx
                })
                
                # Итого (будет обновляться) - жирным шрифтом
                total_label = tk.Label(row_frame, text=f"{result['total_area_m2']:.4f}", 
                                      width=self.col_widths['total'], anchor='center', 
                                      font=('Arial', 10, 'bold'),
                                      fg='#0D47A1',
                                      bg=bg_color)
                total_label.grid(row=0, column=6, padx=0)
                
                # Сохраняем ссылку на label для обновления
                result['total_label'] = total_label
        
        self.update_totals()
        self.status_var.set(f"✓ Загружено DXF файлов: {len(self.files_data)}")
        
        if len(self.files_data) > 0:
            messagebox.showinfo("Загрузка завершена", 
                              f"Найдено и обработано {len(self.files_data)} DXF файлов.\n\n"
                              f"Укажите количество каждой детали в желтых полях.")
    
    def on_quantity_changed(self, index: int):
        """Обработчик изменения количества"""
        try:
            # Получаем новое значение
            if index >= len(self.quantity_entries) or index >= len(self.files_data):
                return
            
            qty_var = self.quantity_entries[index]['var']
            new_qty_str = qty_var.get().strip()
            
            if not new_qty_str:
                return
            
            try:
                new_qty = int(new_qty_str)
                if new_qty < 1 or new_qty > 999:
                    return
            except ValueError:
                return
            
            # Обновляем данные
            data = self.files_data[index]
            old_qty = data.get('quantity', 1)
            
            # Только если количество действительно изменилось
            if new_qty != old_qty:
                data['quantity'] = new_qty
                data['total_area_m2'] = data['area_m2'] * new_qty
                data['total_area_gap_m2'] = data['area_gap_m2'] * new_qty
                
                # Обновляем label итого для этой строки
                if 'total_label' in data:
                    data['total_label'].config(text=f"{data['total_area_m2']:.4f}")
                
                # Обновляем общие итоги
                self.root.after(100, self.update_totals)  # Небольшая задержка для плавности
            
        except Exception as e:
            print(f"Ошибка обновления: {e}")  # Для отладки
    
    def recalculate_totals(self):
        """Пересчет всех итогов"""
        for idx, entry_info in enumerate(self.quantity_entries):
            self.on_quantity_changed(idx)
        
        self.update_totals()
    
    def set_all_quantities(self):
        """Установить количество для всех"""
        qty = simpledialog.askinteger(
            "Количество для всех",
            "Установить одинаковое количество для всех деталей:",
            initialvalue=1,
            minvalue=1,
            maxvalue=999
        )
        
        if qty:
            # Устанавливаем значение во все поля ввода
            for entry_info in self.quantity_entries:
                entry_info['var'].set(str(qty))
            
            # Пересчет произойдет автоматически через trace
    
    def update_totals(self):
        """Обновление итогов"""
        if not self.files_data:
            self.total_label.config(text="")
            return
        
        total_area = sum(d['total_area_m2'] for d in self.files_data)
        total_area_gap = sum(d['total_area_gap_m2'] for d in self.files_data)
        increase = ((total_area_gap / total_area) - 1) * 100 if total_area > 0 else 0
        
        text = f"  ИТОГО: {total_area:.4f} м²  |  С зазорами для листов: {total_area_gap:.4f} м² (+{increase:.1f}%)  "
        self.total_label.config(text=text)
    
    def export_pdf(self):
        """Экспорт в PDF"""
        if not self.files_data:
            messagebox.showwarning("Внимание", "Сначала загрузите файлы!")
            return
        
        # Генерируем PDF
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        
        # Сохранить как
        default_name = f"Расчет_площадей_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
        
        output_path = filedialog.asksaveasfilename(
            title="Сохранить PDF",
            defaultextension=".pdf",
            initialfile=default_name,
            filetypes=[("PDF", "*.pdf")]
        )
        
        if not output_path:
            return
        
        try:
            # Русский шрифт - КРИТИЧНО для кириллицы!
            import os
            try:
                # Пробуем Arial из Windows
                arial_path = r'C:\Windows\Fonts\arial.ttf'
                arial_bold_path = r'C:\Windows\Fonts\arialbd.ttf'
                
                if os.path.exists(arial_path):
                    pdfmetrics.registerFont(TTFont('Arial', arial_path))
                    if os.path.exists(arial_bold_path):
                        pdfmetrics.registerFont(TTFont('Arial-Bold', arial_bold_path))
                    font_name = 'Arial'
                    font_name_bold = 'Arial-Bold'
                else:
                    # Fallback на Times
                    times_path = r'C:\Windows\Fonts\times.ttf'
                    if os.path.exists(times_path):
                        pdfmetrics.registerFont(TTFont('Times', times_path))
                        font_name = 'Times'
                        font_name_bold = 'Times'
                    else:
                        font_name = 'Helvetica'
                        font_name_bold = 'Helvetica-Bold'
            except Exception as e:
                font_name = 'Helvetica'
                font_name_bold = 'Helvetica-Bold'
            
            pdf = SimpleDocTemplate(output_path, pagesize=A4, 
                                   rightMargin=15*mm, leftMargin=15*mm,
                                   topMargin=15*mm, bottomMargin=15*mm)
            
            elements = []
            styles = getSampleStyleSheet()
            
            # Заголовок - с использованием нашего шрифта
            from reportlab.lib.styles import ParagraphStyle
            
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Title'],
                fontName=font_name_bold if 'font_name_bold' in locals() else font_name,
                fontSize=18,
                alignment=1
            )
            
            title = Paragraph("Расчет площадей разверток", title_style)
            elements.append(title)
            elements.append(Spacer(1, 10*mm))
            
            # Информация - с нашим шрифтом
            normal_style = ParagraphStyle(
                'CustomNormal',
                parent=styles['Normal'],
                fontName=font_name,
                fontSize=10
            )
            
            order = self.order_number.get()
            if order:
                info = Paragraph(f"Номер заказа: <b>{order}</b>", normal_style)
                elements.append(info)
            
            date_text = Paragraph(f"Дата: {datetime.now().strftime('%d.%m.%Y %H:%M')}", normal_style)
            elements.append(date_text)
            elements.append(Spacer(1, 10*mm))
            
            # Таблица
            table_data = [['№', 'Файл', 'Габариты, мм', 'Площадь, м²', 'Кол-во', 'Итого, м²']]
            
            for idx, data in enumerate(self.files_data, 1):
                table_data.append([
                    str(idx),
                    data['filename'][:40],
                    f"{data['width']:.0f} × {data['height']:.0f}",
                    f"{data['area_m2']:.4f}",
                    str(data['quantity']),
                    f"{data['total_area_m2']:.4f}"
                ])
            
            # Итоги
            total_area = sum(d['total_area_m2'] for d in self.files_data)
            total_gap = sum(d['total_area_gap_m2'] for d in self.files_data)
            
            table_data.append(['', 'ИТОГО (площадь деталей):', '', '', '', f"{total_area:.4f}"])
            table_data.append(['', 'С зазорами (для листов):', '', '', '', f"{total_gap:.4f}"])
            
            table = Table(table_data, colWidths=[10*mm, 70*mm, 35*mm, 25*mm, 20*mm, 25*mm])
            table.setStyle(TableStyle([
                # Заголовок
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1976D2')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), font_name_bold if 'font_name_bold' in locals() else font_name),
                ('FONTSIZE', (0, 0), (-1, 0), 9),
                ('TOPPADDING', (0, 0), (-1, 0), 6),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
                # Данные
                ('BACKGROUND', (0, 1), (-1, -3), colors.beige),
                ('FONTNAME', (0, 1), (-1, -3), font_name),
                ('FONTSIZE', (0, 1), (-1, -3), 9),
                ('GRID', (0, 0), (-1, -3), 0.5, colors.grey),
                ('ALIGN', (2, 1), (-1, -1), 'RIGHT'),
                ('VALIGN', (0, 1), (-1, -3), 'MIDDLE'),
                # Итоговые строки
                ('BACKGROUND', (0, -2), (-1, -2), colors.HexColor('#FFC107')),
                ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#FFE082')),
                ('FONTNAME', (0, -2), (-1, -1), font_name_bold if 'font_name_bold' in locals() else font_name),
                ('FONTSIZE', (0, -2), (-1, -2), 10),
                ('FONTSIZE', (0, -1), (-1, -1), 9),
                ('LINEABOVE', (0, -2), (-1, -2), 2, colors.black),
                ('TOPPADDING', (0, -2), (-1, -1), 6),
                ('BOTTOMPADDING', (0, -2), (-1, -1), 6),
            ]))
            
            elements.append(table)
            
            # Генерируем
            pdf.build(elements)
            
            messagebox.showinfo("Успех", f"PDF создан:\n{Path(output_path).name}\n\nРасположение:\n{output_path}")
            self.status_var.set(f"PDF создан: {Path(output_path).name}")
            
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось создать PDF:\n{e}")
    
    def export_excel(self):
        """Экспорт в Excel"""
        if not self.files_data:
            messagebox.showwarning("Внимание", "Сначала загрузите файлы!")
            return
        
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill
        except ImportError:
            messagebox.showerror("Ошибка", "Библиотека openpyxl не установлена!\n\npip install openpyxl")
            return
        
        # Сохранить как
        default_name = f"Расчет_площадей_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        
        output_path = filedialog.asksaveasfilename(
            title="Сохранить Excel",
            defaultextension=".xlsx",
            initialfile=default_name,
            filetypes=[("Excel", "*.xlsx")]
        )
        
        if not output_path:
            return
        
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Площади разверток"
            
            # Заголовок
            ws['A1'] = "Расчет площадей разверток"
            ws['A1'].font = Font(size=16, bold=True)
            
            if self.order_number.get():
                ws['A2'] = f"Номер заказа: {self.order_number.get()}"
                ws['A2'].font = Font(size=12)
            
            ws['A3'] = f"Дата: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
            
            # Шапка таблицы
            headers = ['№', 'Файл', 'Ширина, мм', 'Высота, мм', 'Площадь, м²', 'Кол-во', 'Итого, м²']
            row = 5
            
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row, col, header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color='1976D2', fill_type='solid')
                cell.alignment = Alignment(horizontal='center')
            
            # Данные
            for idx, data in enumerate(self.files_data, 1):
                row += 1
                ws.cell(row, 1, idx)
                ws.cell(row, 2, data['filename'])
                ws.cell(row, 3, f"{data['width']:.1f}")
                ws.cell(row, 4, f"{data['height']:.1f}")
                ws.cell(row, 5, f"{data['area_m2']:.4f}")
                ws.cell(row, 6, data['quantity'])
                ws.cell(row, 7, f"{data['total_area_m2']:.4f}")
            
            # Итоги
            row += 1
            total_area = sum(d['total_area_m2'] for d in self.files_data)
            total_gap = sum(d['total_area_gap_m2'] for d in self.files_data)
            
            ws.cell(row, 2, "ИТОГО (площадь деталей):")
            ws.cell(row, 7, f"{total_area:.4f}")
            ws.cell(row, 2).font = Font(bold=True)
            
            row += 1
            ws.cell(row, 2, "С зазорами (для листов):")
            ws.cell(row, 7, f"{total_gap:.4f}")
            ws.cell(row, 2).font = Font(bold=True)
            
            # Авторазмер колонок
            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                ws.column_dimensions[column].width = max_length + 2
            
            wb.save(output_path)
            
            messagebox.showinfo("Успех", f"Excel создан:\n{Path(output_path).name}")
            self.status_var.set(f"Excel создан: {Path(output_path).name}")
            
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось создать Excel:\n{e}")
    
    def run(self):
        """Запуск GUI"""
        self.root.mainloop()


if __name__ == "__main__":
    app = UnfoldingAreaGUI()
    app.run()

