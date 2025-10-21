#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Веб-интерфейс для просмотра базы обрезков
Запускает локальный веб-сервер для удобного просмотра и поиска
"""

import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pathlib import Path
from http.server import HTTPServer, BaseHTTPRequestHandler
from urllib.parse import parse_qs, urlparse
import json
import openpyxl


class WasteDatabase:
    """Работа с базой данных обрезков"""
    
    def __init__(self, db_path: Path = Path("База_обрезков.xlsx")):
        self.db_path = db_path
    
    def get_all_wastes(self, status_filter: str = None):
        """Получает все обрезки из базы"""
        if not self.db_path.exists():
            return []
        
        wb = openpyxl.load_workbook(self.db_path)
        ws = wb.active
        
        wastes = []
        
        for row in range(2, ws.max_row + 1):
            waste = {
                'id': ws.cell(row, 1).value,
                'project': ws.cell(row, 2).value,
                'order': ws.cell(row, 3).value,
                'date': ws.cell(row, 4).value,
                'sheet_number': ws.cell(row, 5).value,
                'width': ws.cell(row, 6).value,
                'height': ws.cell(row, 7).value,
                'area_m2': ws.cell(row, 8).value,
                'material': ws.cell(row, 9).value,
                'location': ws.cell(row, 10).value,
                'status': ws.cell(row, 11).value,
                'used_in': ws.cell(row, 12).value,
                'recommendations': ws.cell(row, 13).value,
                'note': ws.cell(row, 14).value,
            }
            
            # Фильтр по статусу
            if status_filter and waste['status'] != status_filter:
                continue
            
            wastes.append(waste)
        
        return wastes
    
    def get_statistics(self):
        """Статистика базы"""
        all_wastes = self.get_all_wastes()
        
        stats = {
            'total': len(all_wastes),
            'available': 0,
            'used': 0,
            'disposed': 0,
            'total_area_available': 0,
            'total_value_available': 0
        }
        
        for waste in all_wastes:
            status = waste['status']
            
            if status == 'В наличии':
                stats['available'] += 1
                if waste['area_m2']:
                    stats['total_area_available'] += waste['area_m2']
            elif status == 'Использован':
                stats['used'] += 1
            elif status == 'Утилизирован':
                stats['disposed'] += 1
        
        stats['total_value_available'] = stats['total_area_available'] * 3500
        
        if stats['total'] > 0:
            stats['reuse_percent'] = (stats['used'] / stats['total']) * 100
        else:
            stats['reuse_percent'] = 0
        
        return stats


class WasteViewerHandler(BaseHTTPRequestHandler):
    """HTTP обработчик для веб-интерфейса"""
    
    def do_GET(self):
        """Обработка GET запросов"""
        parsed_path = urlparse(self.path)
        path = parsed_path.path
        
        if path == '/':
            self.serve_html()
        elif path == '/api/wastes':
            self.serve_wastes_api(parsed_path.query)
        elif path == '/api/statistics':
            self.serve_statistics_api()
        else:
            self.send_error(404)
    
    def serve_html(self):
        """Отдает HTML страницу"""
        html = self.generate_html()
        
        self.send_response(200)
        self.send_header('Content-type', 'text/html; charset=utf-8')
        self.end_headers()
        self.wfile.write(html.encode('utf-8'))
    
    def serve_wastes_api(self, query_string):
        """API для получения обрезков"""
        params = parse_qs(query_string)
        status_filter = params.get('status', [None])[0]
        
        db = WasteDatabase()
        wastes = db.get_all_wastes(status_filter)
        
        self.send_response(200)
        self.send_header('Content-type', 'application/json; charset=utf-8')
        self.end_headers()
        self.wfile.write(json.dumps(wastes, ensure_ascii=False).encode('utf-8'))
    
    def serve_statistics_api(self):
        """API для получения статистики"""
        db = WasteDatabase()
        stats = db.get_statistics()
        
        self.send_response(200)
        self.send_header('Content-type', 'application/json; charset=utf-8')
        self.end_headers()
        self.wfile.write(json.dumps(stats, ensure_ascii=False).encode('utf-8'))
    
    def generate_html(self):
        """Генерирует HTML страницу"""
        return '''<!DOCTYPE html>
<html lang="ru">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>База обрезков ZVD</title>
    <style>
        * {
            margin: 0;
            padding: 0;
            box-sizing: border-box;
        }
        
        body {
            font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
            background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
            min-height: 100vh;
            padding: 20px;
        }
        
        .container {
            max-width: 1400px;
            margin: 0 auto;
        }
        
        .header {
            background: white;
            border-radius: 15px;
            padding: 30px;
            margin-bottom: 20px;
            box-shadow: 0 10px 30px rgba(0,0,0,0.2);
        }
        
        .header h1 {
            color: #333;
            margin-bottom: 10px;
            font-size: 2.5em;
        }
        
        .stats {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(200px, 1fr));
            gap: 15px;
            margin-bottom: 20px;
        }
        
        .stat-card {
            background: white;
            border-radius: 10px;
            padding: 20px;
            box-shadow: 0 5px 15px rgba(0,0,0,0.1);
            text-align: center;
        }
        
        .stat-card h3 {
            color: #666;
            font-size: 0.9em;
            margin-bottom: 10px;
            text-transform: uppercase;
        }
        
        .stat-card .value {
            font-size: 2em;
            font-weight: bold;
            color: #667eea;
        }
        
        .filters {
            background: white;
            border-radius: 10px;
            padding: 20px;
            margin-bottom: 20px;
            box-shadow: 0 5px 15px rgba(0,0,0,0.1);
        }
        
        .filters select, .filters input {
            padding: 10px 15px;
            border: 2px solid #ddd;
            border-radius: 5px;
            font-size: 1em;
            margin-right: 10px;
        }
        
        .filters select:focus, .filters input:focus {
            outline: none;
            border-color: #667eea;
        }
        
        .table-container {
            background: white;
            border-radius: 10px;
            padding: 20px;
            box-shadow: 0 5px 15px rgba(0,0,0,0.1);
            overflow-x: auto;
        }
        
        table {
            width: 100%;
            border-collapse: collapse;
        }
        
        thead {
            background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
            color: white;
        }
        
        th, td {
            padding: 12px;
            text-align: left;
            border-bottom: 1px solid #ddd;
        }
        
        tbody tr:hover {
            background: #f5f5f5;
            cursor: pointer;
        }
        
        .badge {
            display: inline-block;
            padding: 5px 10px;
            border-radius: 20px;
            font-size: 0.85em;
            font-weight: bold;
        }
        
        .badge-available {
            background: #4CAF50;
            color: white;
        }
        
        .badge-used {
            background: #2196F3;
            color: white;
        }
        
        .badge-disposed {
            background: #F44336;
            color: white;
        }
        
        .loading {
            text-align: center;
            padding: 50px;
            font-size: 1.2em;
            color: #666;
        }
        
        .empty {
            text-align: center;
            padding: 50px;
            color: #999;
        }
        
        .search-highlight {
            background: yellow;
            font-weight: bold;
        }
    </style>
</head>
<body>
    <div class="container">
        <div class="header">
            <h1>🏭 База обрезков ZVD GROUP</h1>
            <p style="color: #666; margin-top: 10px;">Система учета и переиспользования обрезков металла</p>
        </div>
        
        <div class="stats" id="stats">
            <div class="stat-card">
                <h3>Всего обрезков</h3>
                <div class="value" id="stat-total">-</div>
            </div>
            <div class="stat-card">
                <h3>📦 В наличии</h3>
                <div class="value" id="stat-available">-</div>
            </div>
            <div class="stat-card">
                <h3>✓ Использовано</h3>
                <div class="value" id="stat-used">-</div>
            </div>
            <div class="stat-card">
                <h3>📐 Площадь (м²)</h3>
                <div class="value" id="stat-area">-</div>
            </div>
            <div class="stat-card">
                <h3>💰 Стоимость</h3>
                <div class="value" id="stat-value">-</div>
            </div>
            <div class="stat-card">
                <h3>♻️ Переиспользование</h3>
                <div class="value" id="stat-reuse">-</div>
            </div>
        </div>
        
        <div class="filters">
            <select id="filter-status" onchange="loadWastes()">
                <option value="">Все статусы</option>
                <option value="В наличии">В наличии</option>
                <option value="Использован">Использован</option>
                <option value="Утилизирован">Утилизирован</option>
            </select>
            
            <input type="text" id="search" placeholder="🔍 Поиск по ID, проекту, заказу..." 
                   onkeyup="filterTable()" style="width: 300px;">
        </div>
        
        <div class="table-container">
            <table id="wastes-table">
                <thead>
                    <tr>
                        <th>ID</th>
                        <th>Проект</th>
                        <th>Заказ</th>
                        <th>Дата</th>
                        <th>Размеры (мм)</th>
                        <th>Площадь (м²)</th>
                        <th>Материал</th>
                        <th>Расположение</th>
                        <th>Статус</th>
                        <th>Рекомендации</th>
                    </tr>
                </thead>
                <tbody id="wastes-tbody">
                    <tr>
                        <td colspan="10" class="loading">Загрузка данных...</td>
                    </tr>
                </tbody>
            </table>
        </div>
    </div>
    
    <script>
        // Загрузка статистики
        function loadStatistics() {
            fetch('/api/statistics')
                .then(response => response.json())
                .then(stats => {
                    document.getElementById('stat-total').textContent = stats.total;
                    document.getElementById('stat-available').textContent = stats.available;
                    document.getElementById('stat-used').textContent = stats.used;
                    document.getElementById('stat-area').textContent = stats.total_area_available.toFixed(2);
                    document.getElementById('stat-value').textContent = 
                        stats.total_value_available.toLocaleString('ru-RU') + ' ₽';
                    document.getElementById('stat-reuse').textContent = 
                        stats.reuse_percent.toFixed(1) + '%';
                });
        }
        
        // Загрузка обрезков
        function loadWastes() {
            const status = document.getElementById('filter-status').value;
            const url = '/api/wastes' + (status ? '?status=' + encodeURIComponent(status) : '');
            
            fetch(url)
                .then(response => response.json())
                .then(wastes => {
                    displayWastes(wastes);
                });
        }
        
        // Отображение обрезков в таблице
        function displayWastes(wastes) {
            const tbody = document.getElementById('wastes-tbody');
            
            if (wastes.length === 0) {
                tbody.innerHTML = '<tr><td colspan="10" class="empty">Обрезки не найдены</td></tr>';
                return;
            }
            
            let html = '';
            
            for (const waste of wastes) {
                const statusClass = 
                    waste.status === 'В наличии' ? 'badge-available' :
                    waste.status === 'Использован' ? 'badge-used' :
                    'badge-disposed';
                
                html += `
                    <tr>
                        <td><strong>${waste.id}</strong></td>
                        <td>${waste.project || '-'}</td>
                        <td>${waste.order || '-'}</td>
                        <td>${waste.date || '-'}</td>
                        <td>${waste.width} × ${waste.height}</td>
                        <td>${waste.area_m2 ? waste.area_m2.toFixed(4) : '-'}</td>
                        <td>${waste.material || '-'}</td>
                        <td>${waste.location || '-'}</td>
                        <td><span class="badge ${statusClass}">${waste.status}</span></td>
                        <td>${waste.recommendations || '-'}</td>
                    </tr>
                `;
            }
            
            tbody.innerHTML = html;
        }
        
        // Поиск в таблице
        function filterTable() {
            const searchValue = document.getElementById('search').value.toLowerCase();
            const rows = document.querySelectorAll('#wastes-tbody tr');
            
            rows.forEach(row => {
                const text = row.textContent.toLowerCase();
                row.style.display = text.includes(searchValue) ? '' : 'none';
            });
        }
        
        // Инициализация
        loadStatistics();
        loadWastes();
        
        // Автообновление каждые 30 секунд
        setInterval(() => {
            loadStatistics();
            loadWastes();
        }, 30000);
    </script>
</body>
</html>'''
    
    def log_message(self, format, *args):
        """Отключаем логирование запросов"""
        pass


def start_server(port: int = 8080):
    """Запускает веб-сервер"""
    
    # Проверяем наличие базы
    if not Path("База_обрезков.xlsx").exists():
        print("⚠️  База обрезков не найдена!")
        print("Создайте базу, импортировав раскрой из CypCut")
        return
    
    server_address = ('', port)
    httpd = HTTPServer(server_address, WasteViewerHandler)
    
    print("=" * 80)
    print("🌐 ВЕБ-ИНТЕРФЕЙС БАЗЫ ОБРЕЗКОВ")
    print("=" * 80)
    print(f"\n✅ Сервер запущен!")
    print(f"\n🔗 Откройте в браузере:")
    print(f"   http://localhost:{port}")
    print(f"   или")
    print(f"   http://127.0.0.1:{port}")
    print(f"\n💡 Для остановки нажмите Ctrl+C")
    print("\n" + "=" * 80)
    
    # Открываем браузер
    import webbrowser
    import threading
    
    def open_browser():
        import time
        time.sleep(1)
        webbrowser.open(f'http://localhost:{port}')
    
    threading.Thread(target=open_browser, daemon=True).start()
    
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        print("\n\n🛑 Сервер остановлен")
        httpd.server_close()


if __name__ == '__main__':
    start_server(port=8080)


